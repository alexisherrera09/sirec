"""Constructor del documento de la Entrega 2 sobre la plantilla oficial de UNIR.

Reglas de formato que impone la plantilla y el documento de instrucciones
(`instrucciones.pdf`, apartado 1.3) y que este modulo respeta:

- Los estilos Titulo 1/2/3 ya traen numeracion automatica: los titulos se
  escriben SIN numero. Escribir "1. Introduccion" produciria "1. 1. Introduccion".
- Titulo de tablas y figuras ARRIBA: "Tabla N. " en negrita + nombre en cursiva.
- Fuente de tablas y figuras ABAJO, con el estilo "Pie de foto-tabla".
- Cuerpo de tablas: se puede reducir hasta 9 pt cuando hay mucha informacion.
- Cada capitulo arranca en pagina nueva.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PLANTILLA = "plantilla.docx"

# --- estilos de la plantilla (nombres tal como los expone python-docx) ---
S_NORMAL = "Normal"
S_H1 = "Heading 1"
S_H2 = "Heading 2"
S_H3 = "Heading 3"
S_H1_SIN_NUM = "Título 1 sin numerar"
S_REFS = "Referencias bibliográficas"
S_ANEXO = "Anexo"
S_PIE = "Pie de foto-tabla"
S_FIGURA = "Figuras"
S_LISTA = "List Paragraph"
S_PORTADA = "No Spacing"
S_TITULO_INDICE = "Título Índices"
S_TABLA = "Table Grid"


class Documento:
    """Envuelve el .docx y lleva la cuenta de tablas y figuras."""

    def __init__(self, plantilla=PLANTILLA):
        self.doc = Document(plantilla)
        self.n_tabla = 0
        self.n_figura = 0
        self._preparar()

    # ------------------------------------------------------------------ setup
    def _preparar(self):
        """Deja la plantilla con portada + indice y sin el contenido de ejemplo."""
        body = self.doc.element.body
        hijos = list(body.iterchildren())

        # El contenido de ejemplo va del primer 'Heading 1' hasta antes de sectPr.
        primer_h1 = None
        for i, ch in enumerate(hijos):
            if ch.tag == qn("w:p"):
                estilo = ch.find(qn("w:pPr") + "/" + qn("w:pStyle"))
                if estilo is not None and estilo.get(qn("w:val")) == "Ttulo1":
                    primer_h1 = i
                    break
        if primer_h1 is None:
            raise RuntimeError("No se encontro el primer Titulo 1 en la plantilla")

        for ch in hijos[primer_h1:]:
            if ch.tag != qn("w:sectPr"):
                body.remove(ch)

        self._marcador_indice(hijos)
        self._actualizar_campos_al_abrir()
        self._numeracion_jerarquica()

    def _numeracion_jerarquica(self):
        """Numera el 2.o nivel de titulos como 1.1 y el 3.o como 1.1.1.

        La plantilla trae el nivel 2 con vineta (Wingdings) en lugar de numero,
        pero el profesor exige numeracion jerarquica en el cuerpo y en el indice
        (2.1 Objetivo general, 3.1 Empatizar, 3.1.1 ...). Se cambia unicamente la
        definicion de la numeracion; la tipografia, el color y el tamano siguen
        viniendo de los estilos Titulo 1/2/3 de la plantilla.
        """
        try:
            numbering = self.doc.part.numbering_part.element
        except (AttributeError, KeyError):
            return

        objetivo = {0: "%1.", 1: "%1.%2", 2: "%1.%2.%3"}
        for abstract in numbering.findall(qn("w:abstractNum")):
            estilos = {
                lvl.get(qn("w:ilvl")): lvl
                for lvl in abstract.findall(qn("w:lvl"))
            }
            # solo la lista que gobierna los titulos
            pstyles = {
                lvl.find(qn("w:pStyle")).get(qn("w:val"))
                for lvl in estilos.values()
                if lvl.find(qn("w:pStyle")) is not None
            }
            if not {"Ttulo1", "Ttulo2"} <= pstyles:
                continue

            for ilvl, texto in objetivo.items():
                lvl = estilos.get(str(ilvl))
                if lvl is None:
                    continue
                fmt = lvl.find(qn("w:numFmt"))
                if fmt is not None:
                    fmt.set(qn("w:val"), "decimal")
                lvl_text = lvl.find(qn("w:lvlText"))
                if lvl_text is not None:
                    lvl_text.set(qn("w:val"), texto)
                start = lvl.find(qn("w:start"))
                if start is not None:
                    start.set(qn("w:val"), "1")
                # separador: la tabulacion colapsa en los titulos, se usa espacio
                suff = lvl.find(qn("w:suff"))
                if suff is None:
                    suff = OxmlElement("w:suff")
                    lvl_text.addnext(suff)
                suff.set(qn("w:val"), "space")
                # la vineta traia fuente de simbolos; el numero debe heredar la del titulo
                rpr = lvl.find(qn("w:rPr"))
                if rpr is not None:
                    fuentes = rpr.find(qn("w:rFonts"))
                    if fuentes is not None:
                        rpr.remove(fuentes)

    def _marcador_indice(self, hijos):
        """Sustituye el indice de ejemplo por un campo TOC que Word recalcula."""
        body = self.doc.element.body
        tocs = []
        for ch in hijos:
            if ch.tag == qn("w:p"):
                estilo = ch.find(qn("w:pPr") + "/" + qn("w:pStyle"))
                if estilo is not None and estilo.get(qn("w:val"), "").startswith("TDC"):
                    tocs.append(ch)
        if not tocs:  # nombres alternos del estilo de indice
            for ch in hijos:
                if ch.tag == qn("w:p"):
                    estilo = ch.find(qn("w:pPr") + "/" + qn("w:pStyle"))
                    val = estilo.get(qn("w:val"), "") if estilo is not None else ""
                    if val.lower().startswith("toc"):
                        tocs.append(ch)
        if not tocs:
            return

        ancla = tocs[0]
        p = OxmlElement("w:p")
        ancla.addprevious(p)

        # { TOC \o "1-3" \h \z \u }
        r1 = OxmlElement("w:r")
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "begin")
        r1.append(fld)
        r2 = OxmlElement("w:r")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = r' TOC \o "1-3" \h \z \u '
        r2.append(instr)
        r3 = OxmlElement("w:r")
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        r3.append(sep)
        r4 = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "Actualice el índice con clic derecho › Actualizar campos (F9)."
        r4.append(t)
        r5 = OxmlElement("w:r")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        r5.append(end)
        for r in (r1, r2, r3, r4, r5):
            p.append(r)

        for ch in tocs:
            ch.getparent().remove(ch)

    def _actualizar_campos_al_abrir(self):
        """updateFields=true: Word regenera el indice al abrir el documento."""
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            el = OxmlElement("w:updateFields")
            el.set(qn("w:val"), "true")
            settings.append(el)

    # ---------------------------------------------------------------- portada
    def portada(self, titulo, integrantes, fecha):
        """Rellena la portada de la plantilla: titulo, integrantes y fecha.

        No se anade nada mas: cualquier parrafo extra desplaza la tabla y parte
        la portada en dos paginas.
        """
        parrafos = [p for p in self.doc.paragraphs if p.style.name == S_PORTADA]
        textos = [
            "Universidad Internacional de La Rioja",
            "Escuela Superior de Ingeniería y Tecnología",
            "Maestría en Inteligencia Artificial",
            titulo,
        ]
        i = 0
        for p in parrafos:
            if p.text.strip():
                if i < len(textos):
                    self._sustituir_texto(p, textos[i])
                    i += 1

        # tabla de la portada: los rotulos "Trabajo de innovacion presentado por:" y
        # "Fecha:" ya vienen de la plantilla y no se tocan; solo se rellenan los valores.
        if self.doc.tables:
            t = self.doc.tables[0]
            celda = t.cell(0, 1)
            celda.text = ""
            # solo los integrantes: si la celda crece, la fila de la fecha se
            # desborda a la pagina siguiente y la portada queda partida
            for n, nombre in enumerate(integrantes):
                par = celda.paragraphs[0] if n == 0 else celda.add_paragraph()
                par.text = nombre
            t.cell(1, 1).text = fecha

    @staticmethod
    def _sustituir_texto(parrafo, texto):
        if parrafo.runs:
            parrafo.runs[0].text = texto
            for r in parrafo.runs[1:]:
                r.text = ""
        else:
            parrafo.add_run(texto)

    def encabezado(self, autores, titulo_tfe):
        """Encabezado exigido: nombre del estudiante + titulo del trabajo."""
        for sec in self.doc.sections:
            parrafos = sec.header.paragraphs
            if len(parrafos) >= 1:
                self._sustituir_texto(parrafos[0], autores)
            if len(parrafos) >= 2:
                self._sustituir_texto(parrafos[1], titulo_tfe)

    # --------------------------------------------------------------- contenido
    def h1(self, texto, salto=True):
        p = self.doc.add_paragraph(texto, style=S_H1)
        if salto:
            p.paragraph_format.page_break_before = True
        return p

    def h2(self, texto):
        return self.doc.add_paragraph(texto, style=S_H2)

    def h3(self, texto):
        return self.doc.add_paragraph(texto, style=S_H3)

    def h2_sin_numero(self, texto):
        """Subtitulo de 2.o nivel sin numeracion automatica (para los anexos,
        que llevan su propia letra: A.1, A.2...)."""
        p = self.doc.add_paragraph(texto, style=S_H2)
        pPr = p._p.get_or_add_pPr()
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        numId = OxmlElement("w:numId")
        numId.set(qn("w:val"), "0")
        numPr.append(ilvl)
        numPr.append(numId)
        pPr.append(numPr)
        return p

    def h1_sin_numero(self, texto, salto=True):
        p = self.doc.add_paragraph(texto, style=S_H1_SIN_NUM)
        if salto:
            p.paragraph_format.page_break_before = True
        return p

    def anexo(self, texto):
        p = self.doc.add_paragraph(texto, style=S_ANEXO)
        p.paragraph_format.page_break_before = True
        return p

    def p(self, texto, negritas=None):
        """Parrafo de cuerpo. `negritas` es una lista de subcadenas a resaltar."""
        par = self.doc.add_paragraph(style=S_NORMAL)
        if not negritas:
            par.add_run(texto)
            return par
        resto = texto
        for marca in negritas:
            antes, _, resto = resto.partition(marca)
            if antes:
                par.add_run(antes)
            par.add_run(marca).bold = True
        if resto:
            par.add_run(resto)
        return par

    def vineta(self, textos):
        for t in textos:
            par = self.doc.add_paragraph(style=S_LISTA)
            par.style = self.doc.styles[S_LISTA]
            par.paragraph_format.left_indent = Cm(1)
            par.add_run("• " + t)

    def numerada(self, textos):
        for n, t in enumerate(textos, 1):
            par = self.doc.add_paragraph(style=S_LISTA)
            par.paragraph_format.left_indent = Cm(1)
            par.add_run(f"{n}. ").bold = True
            par.add_run(t)

    def referencia(self, texto):
        return self.doc.add_paragraph(texto, style=S_REFS)

    # ----------------------------------------------------------- tablas/figuras
    def rotulo(self, clase, numero, titulo):
        """'Tabla N. ' en negrita + nombre en cursiva, justificado, arriba."""
        par = self.doc.add_paragraph(style=S_NORMAL)
        par.paragraph_format.space_before = Pt(12)
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(f"{clase} {numero}. ")
        r.bold = True
        r2 = par.add_run(titulo if titulo.endswith(".") else titulo + ".")
        r2.italic = True
        return par

    def fuente(self, texto):
        par = self.doc.add_paragraph(style=S_PIE)
        par.add_run(texto)
        return par

    def tabla(self, titulo, encabezados, filas, fuente="Elaboración propia.",
              pt_cuerpo=9.5, anchos=None):
        """Inserta una tabla con rotulo arriba y fuente abajo. Devuelve su numero."""
        self.n_tabla += 1
        self.rotulo("Tabla", self.n_tabla, titulo)

        t = self.doc.add_table(rows=1, cols=len(encabezados))
        t.style = self.doc.styles[S_TABLA]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True

        for i, texto in enumerate(encabezados):
            celda = t.rows[0].cells[i]
            celda.text = ""
            par = celda.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(str(texto))
            run.bold = True
            run.font.size = Pt(pt_cuerpo)
        self._sombrear_fila(t.rows[0], "D9E9F5")

        for fila in filas:
            celdas = t.add_row().cells
            for i, texto in enumerate(fila):
                celdas[i].text = ""
                par = celdas[i].paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = par.add_run(str(texto))
                run.font.size = Pt(pt_cuerpo)

        if anchos:
            for fila in t.rows:
                for i, ancho in enumerate(anchos):
                    fila.cells[i].width = Cm(ancho)

        self.fuente(f"Fuente: {fuente}")
        return self.n_tabla

    @staticmethod
    def _sombrear_fila(fila, hex_color):
        for celda in fila.cells:
            tcPr = celda._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), hex_color)
            tcPr.append(shd)

    def figura(self, titulo, ruta_imagen, fuente="Elaboración propia.", ancho_cm=15.0):
        self.n_figura += 1
        self.rotulo("Figura", self.n_figura, titulo)
        par = self.doc.add_paragraph(style=S_FIGURA)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(ruta_imagen, width=Cm(ancho_cm))
        self.fuente(f"Fuente: {fuente}")
        return self.n_figura

    # -------------------------------------------------------------------- save
    def guardar(self, ruta):
        self.doc.save(ruta)
        return ruta
